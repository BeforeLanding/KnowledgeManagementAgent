import csv
import io
import os
import re
import tempfile
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md", ".eml"}


class NeedsManualProcessing(Exception):
    pass


class PermanentParseError(Exception):
    pass


@dataclass
class Segment:
    locator: str
    text: str


def _clean(text: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", text.replace("\x00", "")).strip()


def parse_bytes(filename: str, content: bytes) -> list[Segment]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise PermanentParseError(f"Unsupported file type: {suffix}")
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
        handle.write(content)
        temporary_path = Path(handle.name)
    try:
        return parse_path(temporary_path, filename)
    finally:
        os.unlink(temporary_path)


def parse_path(path: Path, display_name: str | None = None) -> list[Segment]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        segments = [
            Segment(f"page {i}", _clean(page.extract_text() or ""))
            for i, page in enumerate(PdfReader(path).pages, 1)
        ]
        segments = [item for item in segments if item.text]
        if not segments:
            raise NeedsManualProcessing("PDF has no extractable text layer")
        return segments
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        docx_segments = [
            Segment(f"paragraph {i}", _clean(p.text))
            for i, p in enumerate(doc.paragraphs, 1)
            if _clean(p.text)
        ]
        for table_i, table in enumerate(doc.tables, 1):
            for row_i, row in enumerate(table.rows, 1):
                text = " | ".join(_clean(cell.text) for cell in row.cells)
                if text.strip(" |"):
                    docx_segments.append(Segment(f"table {table_i} row {row_i}", text))
        return docx_segments
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        xlsx_segments: list[Segment] = []
        for sheet in workbook.worksheets:
            for row_i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                text = " | ".join("" if value is None else str(value) for value in row)
                if text.strip(" |"):
                    xlsx_segments.append(Segment(f"sheet {sheet.title} row {row_i}", text))
        return xlsx_segments
    if suffix == ".csv":
        decoded = path.read_text(encoding="utf-8-sig", errors="replace")
        return [
            Segment(f"row {i}", " | ".join(row))
            for i, row in enumerate(csv.reader(io.StringIO(decoded)), 1)
            if any(value.strip() for value in row)
        ]
    if suffix in {".txt", ".md"}:
        lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
        return [
            Segment(f"line {i}", _clean(line)) for i, line in enumerate(lines, 1) if _clean(line)
        ]
    if suffix == ".eml":
        return _parse_email(path.read_bytes())
    raise PermanentParseError(f"Unsupported file type: {suffix}")


def _parse_email(content: bytes) -> list[Segment]:
    message = BytesParser(policy=policy.default).parsebytes(content)
    result = [Segment("email subject", _clean(str(message.get("subject", ""))))]
    for part in message.walk():
        content_disposition = part.get_content_disposition()
        filename = part.get_filename()
        if content_disposition == "attachment" and filename:
            suffix = Path(filename).suffix.lower()
            if suffix in SUPPORTED and suffix != ".eml":
                raw_payload = part.get_payload(decode=True)
                payload = raw_payload if isinstance(raw_payload, bytes) else b""
                for segment in parse_bytes(filename, payload):
                    result.append(
                        Segment(f"attachment {filename}: {segment.locator}", segment.text)
                    )
        elif part.get_content_type() == "text/plain" and content_disposition != "attachment":
            try:
                text = part.get_content()
            except Exception:
                raw_payload = part.get_payload(decode=True)
                text = (
                    raw_payload.decode("utf-8", errors="replace")
                    if isinstance(raw_payload, bytes)
                    else str(raw_payload or "")
                )
            if _clean(text):
                result.append(Segment("email body", _clean(text)))
    return [item for item in result if item.text]


def chunk_segments(
    segments: list[Segment], max_words: int = 700, overlap: int = 100
) -> list[Segment]:
    chunks: list[Segment] = []
    for segment in segments:
        words = segment.text.split()
        if len(words) <= max_words:
            chunks.append(segment)
            continue
        start = 0
        part = 1
        while start < len(words):
            text = " ".join(words[start : start + max_words])
            chunks.append(Segment(f"{segment.locator} part {part}", text))
            if start + max_words >= len(words):
                break
            start += max_words - overlap
            part += 1
    return chunks
