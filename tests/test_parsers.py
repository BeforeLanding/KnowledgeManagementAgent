from email.message import EmailMessage
from io import BytesIO
from types import SimpleNamespace

import pytest
from app.parsers import (
    NeedsManualProcessing,
    PermanentParseError,
    chunk_segments,
    parse_bytes,
)
from docx import Document as DocxDocument
from openpyxl import Workbook


def test_markdown_keeps_line_locators():
    segments = parse_bytes("journal.md", b"# Heading\n\nShipment DEMO-1 cleared.\n")
    assert [item.locator for item in segments] == ["line 1", "line 3"]
    assert "DEMO-1" in segments[1].text


def test_email_body_is_parsed_without_loading_external_content():
    message = EmailMessage()
    message["Subject"] = "Status update"
    message.set_content("Package is ready for dispatch.")
    segments = parse_bytes("message.eml", message.as_bytes())
    assert any(item.locator == "email body" for item in segments)
    assert any("Status update" in item.text for item in segments)


def test_csv_keeps_row_locators():
    segments = parse_bytes("status.csv", "编号,状态\nDEMO-1,已放行\n".encode())
    assert [item.locator for item in segments] == ["row 1", "row 2"]
    assert segments[1].text == "DEMO-1 | 已放行"


def test_docx_keeps_paragraph_and_table_locators():
    document = DocxDocument()
    document.add_paragraph("Synthetic release note")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "DEMO-2"
    table.rows[0].cells[1].text = "Ready"
    content = BytesIO()
    document.save(content)

    segments = parse_bytes("release.docx", content.getvalue())

    assert [(item.locator, item.text) for item in segments] == [
        ("paragraph 1", "Synthetic release note"),
        ("table 1 row 1", "DEMO-2 | Ready"),
    ]


def test_xlsx_keeps_sheet_and_row_locators():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "状态"
    sheet.append(["编号", "状态"])
    sheet.append(["DEMO-3", "完成"])
    content = BytesIO()
    workbook.save(content)
    workbook.close()

    segments = parse_bytes("status.xlsx", content.getvalue())

    assert segments[1].locator == "sheet 状态 row 2"
    assert segments[1].text == "DEMO-3 | 完成"


def test_pdf_keeps_page_locators(monkeypatch):
    pages = [
        SimpleNamespace(extract_text=lambda: "Page one"),
        SimpleNamespace(extract_text=lambda: ""),
    ]
    monkeypatch.setattr("app.parsers.PdfReader", lambda _path: SimpleNamespace(pages=pages))

    segments = parse_bytes("brief.pdf", b"synthetic-pdf-placeholder")

    assert [(item.locator, item.text) for item in segments] == [("page 1", "Page one")]


def test_scanned_pdf_requires_manual_processing(monkeypatch):
    pages = [SimpleNamespace(extract_text=lambda: "")]
    monkeypatch.setattr("app.parsers.PdfReader", lambda _path: SimpleNamespace(pages=pages))
    with pytest.raises(NeedsManualProcessing):
        parse_bytes("scan.pdf", b"synthetic-pdf-placeholder")


def test_unsupported_type_is_rejected():
    with pytest.raises(PermanentParseError):
        parse_bytes("macro.xlsm", b"not safe")


def test_chunking_never_crosses_segment_boundary():
    from app.parsers import Segment

    chunks = chunk_segments(
        [Segment("page 1", "one two three four"), Segment("page 2", "five six")],
        max_words=3,
        overlap=1,
    )
    assert len(chunks) == 3
    assert chunks[-1].locator == "page 2"
